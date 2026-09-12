#!/usr/bin/env python3
"""02장에서 쓰이지 않던 삽화를 설명이 맞는 자리에 그림 2.4 로 넣는다.

'지도 표시와 거리 계산, AI 설명 생성, 개인의 자격 판단을 구분하면' 이라고
설명하는 문단 뒤가 코드 작성·설명 생성·자격 판단 삽화가 놓일 자리다.
"""
from __future__ import annotations

import shutil
import struct
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_polish import rewrite  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "최종교재" / "02장_국내외_인공지능_윤리_동향과_법_제도.docx"
FIGDIR = ROOT / "최종교재" / "그림"
ANCHOR = "지도 표시와 거리 계산, AI 설명 생성, 개인의 자격 판단을 구분하면"
CAPTION = "그림 2.4 — 한 서비스 안의 세 가지 인공지능 이용: 코드 작성, 설명 생성, 자격 판단"
BODY_WIDTH = Emu(5486400)

src = FIGDIR / "그림2-2.png"
dst = FIGDIR / "그림2.4.png"
if src.exists():
    shutil.move(str(src), str(dst))
data = dst.read_bytes()
iw, ih = struct.unpack(">II", data[16:24])

d = Document(str(PATH))
anchor = next((p for p in d.paragraphs if ANCHOR in p.text), None)
if anchor is None:
    print("자리를 찾지 못했습니다")
    raise SystemExit(1)

img_p = d.add_paragraph()
img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
img_p.add_run().add_picture(str(dst), width=BODY_WIDTH,
                            height=Emu(int(int(BODY_WIDTH) * ih / iw)))

cap_p = d.add_paragraph(style=d.styles["Caption"])
rewrite(cap_p, CAPTION)

anchor._p.addnext(cap_p._p)
anchor._p.addnext(img_p._p)

bak = PATH.with_suffix(".docx.fig24bak")
if not bak.exists():
    shutil.copy2(PATH, bak)
d.save(str(PATH))
print(f"그림 2.4 삽입 · {iw}x{ih} · {dst.name}")
