#!/usr/bin/env python3
"""캡션만 있고 그림이 빠진 자리에 그림을 넣는다.

    py tools/insert_figure.py "최종교재/01장_....docx" 1.1 "최종교재/그림/그림1-1_....png"
"""
from __future__ import annotations

import re
import shutil
import struct
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu

ROOT = Path(__file__).resolve().parent.parent
BODY_WIDTH = Emu(5486400)        # 본문 폭 6인치 — 다른 장의 그림과 같은 기준


def size(data: bytes, name: str) -> tuple[int, int]:
    if name.lower().endswith(".png"):
        w, h = struct.unpack(">II", data[16:24])
        return w, h
    # jpeg
    i = 2
    while i < len(data):
        if data[i] != 0xFF:
            i += 1
            continue
        m = data[i + 1]
        if m in (0xC0, 0xC1, 0xC2, 0xC3):
            h, w = struct.unpack(">HH", data[i + 5:i + 9])
            return w, h
        i += 2 + struct.unpack(">H", data[i + 2:i + 4])[0]
    return 1600, 900


def main() -> None:
    path = Path(sys.argv[1])
    num = sys.argv[2]
    img = Path(sys.argv[3])
    data = img.read_bytes()
    iw, ih = size(data, img.name)

    d = Document(str(path))
    cap = None
    for p in d.paragraphs:
        if re.match(rf"^\s*그림\s*{re.escape(num)}\b", p.text):
            cap = p
            break
    if cap is None:
        print(f"그림 {num} 캡션 없음")
        return
    prev = cap._p.getprevious()
    if prev is not None and "graphicData" in prev.xml:
        print(f"그림 {num} 은 이미 들어 있습니다")
        return

    holder = d.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = holder.add_run()
    w = BODY_WIDTH
    run.add_picture(str(img), width=w, height=Emu(int(int(w) * ih / iw)))
    cap._p.addprevious(holder._p)

    bak = path.with_suffix(".docx.insbak")
    if not bak.exists():
        shutil.copy2(path, bak)
    d.save(str(path))
    print(f"그림 {num} 삽입 · {iw}x{ih} · {img.name}")


if __name__ == "__main__":
    main()
