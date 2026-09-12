#!/usr/bin/env python3
"""최종교재 docx 의 문단 스타일과 내용을 훑어본다.

    py tools/docx_dump.py "최종교재/06장_Node.js_설치.docx" --limit 40
"""
from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--limit", type=int, default=40)
    ap.add_argument("--offset", type=int, default=0)
    args = ap.parse_args()

    d = Document(args.path)
    ps = d.paragraphs
    print(f"문단 {len(ps)}개 · 표 {len(d.tables)}개 · 인라인 그림 "
          f"{len(d.inline_shapes)}개")
    for i, p in enumerate(ps):
        if i < args.offset:
            continue
        if i >= args.offset + args.limit:
            break
        t = p.text.strip()
        runs = len(p.runs)
        has_img = "<img>" if "graphicData" in p._p.xml else ""
        print(f"{i:4d} [{p.style.name[:18]:18s}] runs={runs:2d} {has_img} {t[:100]}")


if __name__ == "__main__":
    main()
