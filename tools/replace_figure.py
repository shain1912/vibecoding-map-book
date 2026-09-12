#!/usr/bin/env python3
"""docx 안의 그림 하나를 새 이미지로 바꾸고, 그림 폴더에도 같이 넣는다.

    py tools/replace_figure.py "최종교재/04장_....docx" 4.7 그림작업/roadmap.png
"""
from __future__ import annotations

import re
import shutil
import struct
import sys
from pathlib import Path

from docx import Document
from docx.shared import Emu

ROOT = Path(__file__).resolve().parent.parent
FIGDIR = ROOT / "최종교재" / "그림"


def png_size(data: bytes) -> tuple[int, int]:
    w, h = struct.unpack(">II", data[16:24])
    return w, h


def main() -> None:
    docx_path = Path(sys.argv[1])
    num = sys.argv[2]
    img_path = Path(sys.argv[3])
    data = img_path.read_bytes()
    iw, ih = png_size(data)

    d = Document(str(docx_path))
    cap_i = None
    for i, p in enumerate(d.paragraphs):
        if re.match(rf"^\s*그림\s*{re.escape(num)}\b", p.text):
            cap_i = i
            break
    if cap_i is None:
        print(f"그림 {num} 캡션을 찾지 못했습니다")
        return

    # 캡션 바로 앞에서 그림이 든 문단을 찾는다
    target = None
    for j in range(cap_i - 1, max(cap_i - 4, -1), -1):
        if "graphicData" in d.paragraphs[j]._p.xml:
            target = d.paragraphs[j]
            break
    if target is None:
        print(f"그림 {num} 의 이미지 문단을 찾지 못했습니다")
        return

    blip = target._p.xml
    rid = re.search(r'r:embed="([^"]+)"', blip).group(1)
    part = d.part.related_parts[rid]
    part._blob = data

    # 가로 폭은 그대로 두고 세로만 새 이미지 비율에 맞춘다
    for shape in d.inline_shapes:
        if shape._inline.graphic.graphicData.pic.blipFill.blip.embed == rid:
            w = shape.width
            shape.width = w
            shape.height = Emu(int(w * ih / iw))
            break

    bak = docx_path.with_suffix(".docx.figbak")
    if not bak.exists():
        shutil.copy2(docx_path, bak)
    d.save(str(docx_path))

    out = FIGDIR / f"그림{num}.png"
    for old in FIGDIR.glob(f"그림{num}.*"):
        old.unlink()
    shutil.copy2(img_path, out)
    print(f"그림 {num} 교체 · {iw}x{ih} · {out.name}")


if __name__ == "__main__":
    main()
