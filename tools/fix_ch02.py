#!/usr/bin/env python3
"""02장의 그림 번호를 이어지게 맞추고, 빠져 있던 그림을 넣고, 캡션 스타일을 통일한다.

    py tools/fix_ch02.py --dry
    py tools/fix_ch02.py
"""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_polish import rewrite  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "최종교재" / "02장_국내외_인공지능_윤리_동향과_법_제도.docx"
FIGDIR = ROOT / "최종교재" / "그림"

# 문단 번호 -> 새 캡션 글
NEW_TEXT = {
    20: "표 2.1 인공지능 윤리 논의의 주요 흐름 [H1, H2, K1, K2, I1, I2, E1, E2]",
    37: "그림 2.2 2020년 인공지능 윤리기준과 2026년 대한민국 인공지능 윤리원칙의 구성 비교",
    74: "그림 2.3 인공지능 윤리 원칙을 반영한 지도 서비스 설계 예시",
}
TO_CAPTION = [20, 37, 49, 74, 160]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    d = Document(str(PATH))
    for i, text in NEW_TEXT.items():
        p = d.paragraphs[i]
        print(f"{i:4d} {p.text[:60]}\n     -> {text[:60]}")
        if not args.dry:
            rewrite(p, text)
    for i in TO_CAPTION:
        if not args.dry:
            d.paragraphs[i].style = d.styles["Caption"]
    if not args.dry:
        bak = PATH.with_suffix(".docx.ch02bak")
        if not bak.exists():
            shutil.copy2(PATH, bak)
        d.save(str(PATH))

        # 그림 파일 이름을 새 번호에 맞춘다
        for old, new in [("그림2-3.png", "그림2.2.png"), ("그림2-4.png", "그림2.3.png")]:
            src = FIGDIR / old
            if src.exists():
                shutil.move(str(src), str(FIGDIR / new))
        # 캡션 자리에 그림을 끼워 넣는다
        for num in ("2.2", "2.3"):
            subprocess.run([sys.executable, "tools/insert_figure.py", str(PATH), num,
                            str(FIGDIR / f"그림{num}.png")], cwd=ROOT)
    print("끝")


if __name__ == "__main__":
    main()
