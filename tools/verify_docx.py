#!/usr/bin/env python3
"""바꾼 docx 가 원래 구조를 그대로 지키는지 .bak 과 맞춰 본다.

    py tools/verify_docx.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
FINAL = ROOT / "최종교재"


def shape(path: Path) -> dict:
    d = Document(str(path))
    return {
        "paras": len(d.paragraphs),
        "styles": [p.style.name for p in d.paragraphs],
        "imgs": len(d.inline_shapes),
        "tables": len(d.tables),
        "chars": sum(len(p.text) for p in d.paragraphs),
    }


def main() -> None:
    bad = 0
    for bak in sorted(FINAL.glob("*.docx.bak")):
        cur = bak.with_suffix("")          # ...docx.bak -> ...docx
        a, b = shape(bak), shape(cur)
        msgs = []
        if a["paras"] != b["paras"]:
            msgs.append(f"문단 수 {a['paras']} -> {b['paras']}")
        if a["styles"] != b["styles"]:
            msgs.append("문단 스타일 배열이 달라짐")
        if a["imgs"] != b["imgs"]:
            msgs.append(f"그림 {a['imgs']} -> {b['imgs']}")
        if a["tables"] != b["tables"]:
            msgs.append(f"표 {a['tables']} -> {b['tables']}")
        ratio = b["chars"] / max(a["chars"], 1)
        if not (0.82 <= ratio <= 1.25):
            msgs.append(f"글자 수 {a['chars']:,} -> {b['chars']:,} ({ratio:.0%})")
        name = cur.name
        if msgs:
            bad += 1
            print(f"확인 필요 {name}: " + " · ".join(msgs))
        else:
            print(f"OK {name}  글자 {a['chars']:,} -> {b['chars']:,} ({ratio:.0%})")
    print(f"\n문제 {bad}개")


if __name__ == "__main__":
    main()
