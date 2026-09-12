#!/usr/bin/env python3
"""원고에 짝이 없어 옛 초고 말투로 남은 docx 문단만 모아 따로 다듬는다.

    py tools/extra_polish.py extract      # prose/extra.txt 와 목록을 만든다
    py tools/extra_polish.py apply        # prose/replies/extra.txt 를 되꽂는다

윤문한 본문과 말투가 달라 보이는 곳을 없애기 위한 마무리 단계다.
"""
from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_polish import (BULLET, OFFSET, align, final_path, norm,  # noqa: E402
                          fix_numbers, rewrite)
from dola_prep import HEADER_LINES, parse_reply, prose_items  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
PROSE = ROOT / "prose"
MAP = PROSE / "extra_map.json"

# 상자 제목이나 목록 머리처럼 다듬으면 안 되는 줄
SKIP = re.compile(r"^\s*(💡|⚠|🗣|🍎|📌|✅|▶|■|이 장에서 배우는 것|다음 장 예고|마치며)")


def collect() -> list[tuple[str, int, str]]:
    out: list[tuple[str, int, str]] = []
    for n in range(OFFSET, 30):
        ch = f"ch{n - OFFSET:02d}"
        src_md = ROOT / "윤문비교" / "원문" / f"{ch}.md"
        docx_path = final_path(n)
        if not (src_md.exists() and docx_path):
            continue
        src_texts = [t for _, t in prose_items(src_md.read_text(encoding="utf-8"))]
        d = Document(str(docx_path))
        cand = [(i, p) for i, p in enumerate(d.paragraphs)
                if p.style.name == "Normal" and p.text.strip()
                and not BULLET.match(p.text) and len(p.text.strip()) > 12]
        doc_texts = [p.text for _, p in cand]
        m = align(doc_texts, src_texts)
        for di, (pi, p) in enumerate(cand):
            if di in m or SKIP.match(p.text) or len(p.text) < 40:
                continue
            out.append((docx_path.name, pi, p.text))
    return out


def cmd_extract() -> None:
    items = collect()
    body = "\n\n".join(f"[{i}] {t}" for i, (_, _, t) in enumerate(items, 1))
    PROSE.mkdir(parents=True, exist_ok=True)
    (PROSE / "extra.txt").write_text("\n".join(HEADER_LINES) + body, encoding="utf-8")
    MAP.write_text(json.dumps([[f, p] for f, p, _ in items], ensure_ascii=False),
                   encoding="utf-8")
    print(f"짝 없는 문단 {len(items)}개 · {len(body):,}자 -> {PROSE / 'extra.txt'}")


def cmd_apply(dry: bool) -> None:
    items = collect()
    reply = parse_reply((PROSE / "replies" / "extra.txt").read_text(encoding="utf-8"))
    by_file: dict[str, list[tuple[int, str]]] = {}
    kept = 0
    for i, (fname, pi, old) in enumerate(items, 1):
        new = reply.get(i, "").strip()
        if not new:
            continue
        if not (0.8 <= len(new) / max(len(old), 1) <= 2.4):
            kept += 1
            continue
        if new.count("`") != old.count("`"):
            kept += 1
            continue
        by_file.setdefault(fname, []).append((pi, new))

    for fname, jobs in by_file.items():
        path = ROOT / "최종교재" / fname
        d = Document(str(path))
        for pi, new in jobs:
            p = d.paragraphs[pi]
            if norm(p.text) == norm(new):
                continue
            if not dry:
                rewrite(p, fix_numbers(new))
        if not dry:
            bak = path.with_suffix(".docx.exbak")
            if not bak.exists():
                shutil.copy2(path, bak)
            d.save(str(path))
        print(f"{fname}: {len(jobs)}개 다듬음" + (" [미리보기]" if dry else ""))
    print(f"\n지켜 둔 문단 {kept}개")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "extract"
    if cmd == "extract":
        cmd_extract()
    else:
        cmd_apply("--dry" in sys.argv)
