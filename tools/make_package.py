#!/usr/bin/env python3
"""출판사에 그대로 넘길 수 있게 최종 폴더를 정리해 만든다.

    py tools/make_package.py

만들어지는 모양

    출판사제출/
      본문/          00장~29장 docx
      그림/          그림N.M 이름의 원본 이미지
      그림목록.csv   그림 번호 · 파일명 · 캡션 · 수록 장
      제출안내.md    폴더 구성과 표기 규칙
"""
from __future__ import annotations

import csv
import re
import shutil
from collections import defaultdict
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
FINAL = ROOT / "최종교재"
OUT = ROOT / "출판사제출"

CAP = re.compile(r"^\s*그림\s*([\d.\-]+[\d])\s*[—\-–:]?\s*(.*)$")


def collect() -> tuple[list[tuple[str, str, str, str]], dict[str, int]]:
    rows: list[tuple[str, str, str, str]] = []
    figs = {}
    for p in (FINAL / "그림").iterdir():
        if p.is_file():
            m = re.match(r"^그림\s*([\d.\-]+)", p.stem)
            if m:
                figs[m.group(1).replace("-", ".")] = p.name

    stats: dict[str, int] = {}
    for docx in sorted(FINAL.glob("*.docx")):
        d = Document(str(docx))
        stats[docx.name] = sum(len(x.text) for x in d.paragraphs)
        for para in d.paragraphs:
            if para.style.name != "Caption":
                continue
            m = CAP.match(para.text)
            if not m:
                continue
            num = m.group(1)
            rows.append((num, figs.get(num, "(없음)"), m.group(2).strip(), docx.name))
    return rows, stats


GUIDE = """# 출판사 제출 안내

『바이브코딩으로 만드는 나만의 지도 웹』 원고와 그림을 함께 넘깁니다.

## 폴더 구성

| 폴더·파일 | 내용 |
|---|---|
| `본문/` | 장별 원고 {n_doc}개 (00장~29장, docx) |
| `그림/` | 본문에 들어갈 그림 원본 {n_fig}개 |
| `그림목록.csv` | 그림 번호 · 파일명 · 캡션 · 수록 장 |

## 그림 표기 규칙

- 본문에는 `그림 6.1 — 설치 전 node -v 실행 결과`처럼 캡션 문단으로 자리를 잡아 두었습니다.
- 그림 파일 이름은 캡션의 번호와 같습니다. 예: 캡션 `그림 6.1` -> 파일 `그림6.1.png`
- 대부분 2400×1500 픽셀이며, 화면을 그대로 담은 것이라 확대 없이 그대로 쓰시면 됩니다.
- 1장~3장의 개념 그림은 `그림1-1_편익과_윤리적_문제.png`처럼 이름에 설명이 붙어 있습니다.

## 원고 표기 규칙

- 장 제목은 `Title`, 절 제목은 `Heading 1`, 명령어·코드 블록은 `Source Code` 스타일입니다.
- 본문 속 명령어와 파일 이름은 Consolas 8.5pt로 구분해 두었습니다.
- 장 번호와 그림 번호는 최종 번호 하나만 남겨 두었습니다.

## 장별 분량

{table}
"""


def main() -> None:
    rows, stats = collect()

    if OUT.exists():
        shutil.rmtree(OUT)
    (OUT / "본문").mkdir(parents=True)
    (OUT / "그림").mkdir(parents=True)

    for p in sorted(FINAL.glob("*.docx")):
        shutil.copy2(p, OUT / "본문" / p.name)
    for p in (FINAL / "그림").iterdir():
        if p.is_file():
            shutil.copy2(p, OUT / "그림" / p.name)

    with (OUT / "그림목록.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["그림 번호", "파일명", "캡션", "수록 장"])
        for r in rows:
            w.writerow([f"그림 {r[0]}", r[1], r[2], r[3]])

    per = defaultdict(int)
    for name, n in stats.items():
        per[name] = n
    table = "\n".join(["| 장 | 글자 수 |", "|---|---|"]
                      + [f"| {k.replace('.docx','')} | {v:,} |" for k, v in per.items()])

    (OUT / "제출안내.md").write_text(
        GUIDE.format(n_doc=len(stats), n_fig=len(list((OUT / '그림').iterdir())),
                     table=table),
        encoding="utf-8")

    total = sum(stats.values())
    print(f"{OUT} 생성")
    print(f"  본문 {len(stats)}개 · 그림 {len(list((OUT / '그림').iterdir()))}개 "
          f"· 캡션 {len(rows)}개 · 본문 글자 {total:,}자")


if __name__ == "__main__":
    main()
