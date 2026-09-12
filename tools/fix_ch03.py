#!/usr/bin/env python3
"""03장의 그림·표 번호가 겹치던 것을 나오는 순서대로 다시 매긴다.

공저자 원고라 글은 그대로 두고 번호와 캡션 스타일만 손댄다.
그림 파일 이름도 새 번호에 맞춘다.

    py tools/fix_ch03.py --dry
    py tools/fix_ch03.py
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from apply_polish import rewrite  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
PATH = ROOT / "최종교재" / "03장_개인정보_보호와_저작권_및_디지털_자료의_이용.docx"
FIGDIR = ROOT / "최종교재" / "그림"

# 문단 번호 -> (새 번호, 캡션 스타일로 바꿀지)
CAPTIONS = {
    30: ("그림 3.1", True),
    45: ("그림 3.2", False),
    101: ("그림 3.3", False),
    118: ("그림 3.4", False),
    120: ("그림 3.5", False),
    164: ("그림 3.6", False),
    31: ("표 3.1", False),
    48: ("표 3.2", True),
    54: ("표 3.3", False),
    61: ("표 3.4", False),
    80: ("표 3.5", False),
    110: ("표 3.6", False),
    126: ("표 3.7", True),
    142: ("표 3.8", False),
    153: ("표 3.9", False),
    166: ("표 3.10", False),
    181: ("표 3.11", False),
    188: ("표 3.12", False),
}

# 본문에서 가리키는 번호
BODY = {
    102: ("그림 3.2", "그림 3.3"),
    165: ("그림 3.3", "그림 3.6"),
}

# 캡션 번호가 바뀐 만큼 그림 파일 이름도 맞춘다 (옛 이름 -> 새 이름)
RENAME = [
    ("그림3-5.png", "그림3.4.png"),
    ("그림3-6.png", "그림3.5.png"),
    ("그림3.4.png", "그림3.6.png"),
]

HEAD = re.compile(r"^\s*(그림|표)\s*3\.[0-9○]+\s*")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    d = Document(str(PATH))
    for i, (newnum, to_caption) in sorted(CAPTIONS.items()):
        p = d.paragraphs[i]
        rest = HEAD.sub("", p.text).strip()
        new = f"{newnum} {rest}"
        if new == p.text and not to_caption:
            continue
        print(f"{i:4d} {p.text[:52]}\n     -> {new[:52]}")
        if not args.dry:
            rewrite(p, new)
            if to_caption:
                p.style = d.styles["Caption"]

    for i, (old, new) in BODY.items():
        p = d.paragraphs[i]
        if old in p.text:
            t = p.text.replace(old, new, 1)
            print(f"{i:4d} 본문 참조 {old} -> {new}")
            if not args.dry:
                rewrite(p, t)

    if not args.dry:
        bak = PATH.with_suffix(".docx.ch03bak")
        if not bak.exists():
            shutil.copy2(PATH, bak)
        d.save(str(PATH))

        tmp = FIGDIR / "_tmp_ch03"
        tmp.mkdir(exist_ok=True)
        for old, _ in RENAME:
            src = FIGDIR / old
            if src.exists():
                shutil.move(str(src), str(tmp / old))
        for old, new in RENAME:
            src = tmp / old
            if src.exists():
                shutil.move(str(src), str(FIGDIR / new))
        tmp.rmdir()
        print("그림 파일 이름도 새 번호에 맞췄습니다")


if __name__ == "__main__":
    main()
